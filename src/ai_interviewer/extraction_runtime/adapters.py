"""Versioned, content-only text adapters for the isolated extraction worker.

Each adapter reads validated bytes and returns sanitized Unicode text that already
satisfies the D1 source-text content contract (LF newlines, no control/format/surrogate
characters, bounded length). Adapters never touch the filesystem, network, or database;
they are pure functions so they can run inside the isolated child process unmodified.
"""

from __future__ import annotations

import io
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass

import pypdf
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf.errors import PdfReadError

PDF_ADAPTER = "isolated-pdf-parser"
DOCX_ADAPTER = "isolated-docx-parser"
TEXT_ADAPTER = "isolated-text-parser"
PDF_ADAPTER_VERSION = "1"
DOCX_ADAPTER_VERSION = "1"
TEXT_ADAPTER_VERSION = "1"

MAX_OUTPUT_CHARACTERS = 500_000


class ExtractionAdapterError(RuntimeError):
    """Base type for closed, content-free adapter failures."""

    code: str


class UnsupportedInputError(ExtractionAdapterError):
    code = "input_unsupported"


class CorruptInputError(ExtractionAdapterError):
    code = "input_corrupt"


class EncryptedInputError(ExtractionAdapterError):
    code = "input_encrypted"


class EmptyInputError(ExtractionAdapterError):
    code = "input_empty"


class OutputBoundsExceededError(ExtractionAdapterError):
    code = "resource_exceeded"


def sanitize_extracted_text(text: str) -> str:
    """Normalize adapter output to the exact D1 source-text content contract."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    filtered = "".join(
        character
        for character in normalized
        if not (
            (unicodedata.category(character) == "Cc" and character not in {"\n", "\t"})
            or unicodedata.category(character) in {"Cf", "Cs"}
        )
    )
    return filtered.strip("\n").strip()


def _finalize(text: str) -> str:
    sanitized = sanitize_extracted_text(text)
    if not sanitized:
        raise EmptyInputError("no extractable text was found")
    if len(sanitized) > MAX_OUTPUT_CHARACTERS:
        raise OutputBoundsExceededError("extracted text exceeds the bounded output contract")
    return sanitized


def extract_text_txt(content: bytes) -> str:
    if not content:
        raise EmptyInputError("input is empty")
    try:
        decoded = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise CorruptInputError("input is not valid UTF-8 text") from exc
    return _finalize(decoded)


def extract_text_pdf(content: bytes) -> str:
    if not content:
        raise EmptyInputError("input is empty")
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise EncryptedInputError("PDF content is password protected")
        pages = [page.extract_text() or "" for page in reader.pages]
    except EncryptedInputError:
        raise
    except PdfReadError as exc:
        raise CorruptInputError("PDF structure could not be parsed") from exc
    except Exception as exc:
        raise CorruptInputError("PDF content could not be parsed") from exc
    return _finalize("\n".join(pages))


def extract_text_docx(content: bytes) -> str:
    if not content:
        raise EmptyInputError("input is empty")
    try:
        document = DocxDocument(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
    except PackageNotFoundError as exc:
        raise CorruptInputError("DOCX container could not be parsed") from exc
    except Exception as exc:
        raise CorruptInputError("DOCX content could not be parsed") from exc
    return _finalize("\n".join(paragraphs))


_ADAPTERS: dict[tuple[str, str], Callable[[bytes], str]] = {
    (PDF_ADAPTER, PDF_ADAPTER_VERSION): extract_text_pdf,
    (DOCX_ADAPTER, DOCX_ADAPTER_VERSION): extract_text_docx,
    (TEXT_ADAPTER, TEXT_ADAPTER_VERSION): extract_text_txt,
}


@dataclass(frozen=True, slots=True)
class AdapterIdentity:
    adapter: str
    version: str


def resolve_adapter(adapter: str, version: str) -> Callable[[bytes], str]:
    """Look up an exact, separately versioned adapter; unknown pairs are unsupported."""
    resolved = _ADAPTERS.get((adapter, version))
    if resolved is None:
        raise UnsupportedInputError(f"no adapter for {adapter}:{version}")
    return resolved
